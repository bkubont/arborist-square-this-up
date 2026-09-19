#!/usr/bin/env python3
"""Regenerate data/handyman-catalog.json from Brittany's materials workbook.

Source (project store):
  internal/packet/Handyman_Work_Order_Line_Items_materials.xlsx

Usage:
  python3 scripts/regenerate-handyman-catalog.py [path-to-xlsx]
"""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(
    "/cursor/stores/bc-2f74d22b-9eba-43b0-844e-2a636c35e174/internal/packet/"
    "Handyman_Work_Order_Line_Items_materials.xlsx"
)
OUT = ROOT / "data" / "handyman-catalog.json"
DEFAULT_RATE = 55


def slug(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")
    return s[:80] or "item"


def money(v):
    if v is None or v == "":
        return None
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def parse_hours_mid(avg_time):
    """Parse midpoint hours from ranges like '15–20 min', '1.5–3 hrs', '1–2 days'."""
    if avg_time is None:
        return None
    t = str(avg_time).strip().lower()
    if not t:
        return None
    t = re.sub(r"\(.*?\)", "", t).strip()
    t = t.replace("—", "-").replace("–", "-").replace("−", "-")

    day_m = re.search(r"([\d.]+)\s*-\s*([\d.]+)\s*days?", t)
    if day_m:
        a, b = float(day_m.group(1)), float(day_m.group(2))
        return round(((a + b) / 2) * 8, 2)

    day_s = re.search(r"^([\d.]+)\s*days?$", t)
    if day_s:
        return round(float(day_s.group(1)) * 8, 2)

    hr_m = re.search(r"([\d.]+)\s*-\s*([\d.]+)\s*hrs?", t)
    if hr_m:
        return round((float(hr_m.group(1)) + float(hr_m.group(2))) / 2, 2)

    hr_s = re.search(r"^([\d.]+)\s*hrs?$", t)
    if hr_s:
        return round(float(hr_s.group(1)), 2)

    min_m = re.search(r"([\d.]+)\s*-\s*([\d.]+)\s*min", t)
    if min_m:
        return round(((float(min_m.group(1)) + float(min_m.group(2))) / 2) / 60, 2)

    min_s = re.search(r"^([\d.]+)\s*min", t)
    if min_s:
        return round(float(min_s.group(1)) / 60, 2)

    return None


def drop_none(d: dict) -> dict:
    return {k: v for k, v in d.items() if v is not None}


def load_maintenance(wb):
    ws = wb["Maintenance Catalog"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Service / Item / Task")
    out = {}
    for r in rows[hdr_idx + 1 :]:
        if not r or not r[0]:
            continue
        task = str(r[0]).strip()
        hours = money(r[3])
        out[task] = drop_none(
            {
                "hours_mid": hours,
                "tools": str(r[5]).strip() if len(r) > 5 and r[5] else None,
                "notes": str(r[6]).strip() if len(r) > 6 and r[6] else None,
                "suggested_package": str(r[7]).strip() if len(r) > 7 and r[7] else None,
            }
        )
    return out


def load_line_sheet(wb, sheet_name, source_key, maint):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Service / Item / Task")
    items = []
    for r in rows[hdr_idx + 1 :]:
        if not r or not r[0] or not str(r[0]).strip():
            continue
        task = str(r[0]).strip()
        category = str(r[1]).strip() if r[1] else ""
        avg_time = str(r[2]).strip() if r[2] else ""
        labor_rate = money(r[3]) or DEFAULT_RATE
        est_labor = money(r[4])
        est_materials = money(r[5])
        mat_vol = str(r[6]).strip() if r[6] else None
        mat_flag = str(r[7]).strip() if r[7] else None
        ballpark = money(r[8])
        materials_note = str(r[9]).strip() if len(r) > 9 and r[9] else None
        tools = str(r[10]).strip() if len(r) > 10 and r[10] else ""
        maintenance = str(r[11]).strip() if len(r) > 11 and r[11] else "One-off"
        notes = str(r[12]).strip() if len(r) > 12 and r[12] else ""

        m = maint.get(task, {})
        hours_mid = m.get("hours_mid")
        if hours_mid is None:
            hours_mid = parse_hours_mid(avg_time)
        if hours_mid is None and est_labor is not None and labor_rate:
            hours_mid = round(est_labor / labor_rate, 2)

        if not notes and m.get("notes"):
            notes = m["notes"]
        if not tools and m.get("tools"):
            tools = m["tools"]

        if ballpark is None and est_labor is not None:
            ballpark = round(est_labor + (est_materials or 0), 2)

        item = drop_none(
            {
                "id": f"{source_key}-{slug(task)}",
                "task": task,
                "category": category,
                "avg_time": avg_time,
                "hours_mid": hours_mid,
                "labor_rate": labor_rate,
                "est_labor_cost": est_labor,
                "est_materials_cost": est_materials,
                "materials_volatility": mat_vol,
                "materials_flag": mat_flag,
                "ballpark_total": ballpark,
                "materials_note": materials_note,
                "tools": tools or "",
                "maintenance": maintenance or "One-off",
                "notes": notes or "",
                "source": source_key,
                "suggested_package": m.get("suggested_package"),
            }
        )
        items.append(item)
    return items


def load_materials_master(wb):
    ws = wb["Materials_Master"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Item / Material")
    materials = []
    for r in rows[hdr_idx + 1 :]:
        if not r or not r[0]:
            continue
        materials.append(
            drop_none(
                {
                    "id": f"material-{slug(r[0])}",
                    "name": str(r[0]).strip(),
                    "category": str(r[1]).strip() if r[1] else "",
                    "unit": str(r[2]).strip() if r[2] else "",
                    "typical_qty_note": str(r[3]).strip() if r[3] else None,
                    "avg_unit_cost": money(r[4]),
                    "low_typical": money(r[5]),
                    "high_typical": money(r[6]),
                    "volatility": str(r[7]).strip() if r[7] else None,
                    "flag": str(r[8]).strip() if r[8] else None,
                    "common_jobs": str(r[9]).strip() if len(r) > 9 and r[9] else None,
                    "preferred_source": str(r[10]).strip() if len(r) > 10 and r[10] else None,
                    "notes": str(r[11]).strip() if len(r) > 11 and r[11] else None,
                }
            )
        )
    return materials


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        raise SystemExit(f"Workbook not found: {xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    maint = load_maintenance(wb)
    everyday = load_line_sheet(wb, "Work Order Line Items", "everyday", maint)
    less = load_line_sheet(wb, "Less Frequent Items", "less_frequent", maint)
    items = everyday + less
    materials = load_materials_master(wb)

    with_mats = sum(1 for i in items if i.get("est_materials_cost") is not None)
    price_check = sum(1 for i in items if i.get("materials_flag"))

    catalog = {
        "version": 2,
        "source": xlsx.name,
        "default_labor_rate": DEFAULT_RATE,
        "market": "ZIP 49718 / Carp Lake–Petoskey",
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z",
        "counts": {
            "line_items": len(items),
            "everyday": len(everyday),
            "less_frequent": len(less),
            "with_materials_price": with_mats,
            "materials_price_check": price_check,
            "materials_master": len(materials),
        },
        "categories": sorted({i["category"] for i in items if i.get("category")}),
        "items": items,
        "materials": materials,
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(catalog, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUT} ({len(items)} line items, {len(materials)} materials)")


if __name__ == "__main__":
    main()
